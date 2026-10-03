"""
smoke_test.py – round-trip tests for encryptor.py
Run: python smoke_test.py
"""

import io
import openpyxl
import pandas as pd
from encryptor import (
    anonymise, restore,
    anonymise_workbook, restore_workbook,
    mapping_to_json, mapping_from_json,
    wb_to_bytes, wb_to_dataframe,
)

# ── shared test data ────────────────────────────────────────────────────────
COLS = ["Name", "City", "Department"]
df = pd.DataFrame({
    "Name":       ["Alice", "Bob", "Alice", "Charlie"],
    "City":       ["Berlin", "Paris", "Berlin", "Rome"],
    "Department": ["Sales", "HR", "Sales", "Engineering"],
    "Salary":     [50000, 60000, 55000, 70000],
    "Score":      [8.5, 7.0, 9.1, 6.3],
})

print("=== Original ===")
print(df.to_string(index=False))


# ── Test 1: DataFrame API round-trip ────────────────────────────────────────
print("\n─── Test 1: DataFrame API ───")
df_anon, mapping = anonymise(df, COLS)
mapping_json = mapping_to_json(mapping)
mapping_loaded = mapping_from_json(mapping_json)
df_restored = restore(df_anon, mapping_loaded)

print("Anonymised:")
print(df_anon.to_string(index=False))
print("Restored:")
print(df_restored.to_string(index=False))
assert df.equals(df_restored), "DataFrame round-trip FAILED!"
print("✅  DataFrame round-trip OK")


# ── Test 2: Workbook API round-trip ─────────────────────────────────────────
print("\n─── Test 2: Workbook API (style-preserving) ───")

# Build a workbook from the test data (with a simple style applied).
wb_orig = openpyxl.Workbook()
ws = wb_orig.active
ws.append(list(df.columns))

# Apply a background colour to the header to verify it survives round-trip.
from openpyxl.styles import PatternFill
yellow = PatternFill(fill_type="solid", fgColor="FFFF00")
for cell in ws[1]:
    cell.fill = yellow

for row in df.itertuples(index=False):
    ws.append(list(row))

# Anonymise.
wb_anon, wb_mapping = anonymise_workbook(wb_orig, COLS)
wb_mapping_json = mapping_to_json(wb_mapping)

df_anon_wb = wb_to_dataframe(wb_anon)
print("Anonymised (workbook):")
print(df_anon_wb.to_string(index=False))

# Restore.
wb_mapping_loaded = mapping_from_json(wb_mapping_json)
wb_restored = restore_workbook(wb_anon, wb_mapping_loaded)

df_restored_wb = wb_to_dataframe(wb_restored)
print("Restored (workbook):")
print(df_restored_wb.to_string(index=False))

assert df.equals(df_restored_wb), "Workbook round-trip FAILED! Values differ."

# Verify the header fill colour survived.
ws_restored = wb_restored.active
for cell in ws_restored[1]:
    # openpyxl stores colours as ARGB (8 hex chars); strip optional alpha prefix.
    rgb = cell.fill.fgColor.rgb
    assert rgb.endswith("FFFF00"), (
        f"Style lost on column '{cell.value}'! fgColor = {rgb}"
    )

print("✅  Workbook round-trip OK (values AND styles preserved)")
print("\n🎉  All tests passed.")
